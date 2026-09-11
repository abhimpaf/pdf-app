import tempfile
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Emu

_EMU_PER_INCH = 914400


def pdf_to_pptx(input_path: Path, output_path: Path, dpi: int = 150) -> Path:
    zoom = dpi / 72
    matrix = pymupdf.Matrix(zoom, zoom)

    doc = pymupdf.open(input_path)
    try:
        if doc.page_count == 0:
            raise ValueError("PDF has no pages to convert")

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            presentation = Presentation()

            for index, page in enumerate(doc):
                page_width_in = page.rect.width / 72
                page_height_in = page.rect.height / 72

                if index == 0:
                    presentation.slide_width = Emu(int(page_width_in * _EMU_PER_INCH))
                    presentation.slide_height = Emu(int(page_height_in * _EMU_PER_INCH))

                pixmap = page.get_pixmap(matrix=matrix, alpha=False)
                image_path = tmp_path / f"page{index}.jpg"
                pixmap.save(image_path, jpg_quality=90)

                slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank layout
                slide.shapes.add_picture(
                    str(image_path),
                    left=0,
                    top=0,
                    width=presentation.slide_width,
                    height=presentation.slide_height,
                )

            presentation.save(output_path)
    finally:
        doc.close()

    return output_path
