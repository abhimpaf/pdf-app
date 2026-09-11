from pptx import Presentation

from pdfapp.convert_ppt import pdf_to_pptx


def test_pdf_to_pptx_creates_one_slide_per_page(sample_pdf, tmp_path):
    output = tmp_path / "out.pptx"
    result = pdf_to_pptx(sample_pdf, output, dpi=100)

    assert result == output
    assert output.exists()

    presentation = Presentation(output)
    slides = list(presentation.slides)
    assert len(slides) == 3

    first_slide = slides[0]
    assert len(first_slide.shapes) == 1  # one full-bleed picture per slide
